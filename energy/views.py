import csv
from io import BytesIO

from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404, redirect
from .forms import BuildingForm
from .calc import calc_heating_demand
from .models import Building
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart




def dashboard(request):
    building_count = Building.objects.count()
    last_building = Building.objects.order_by("-id").first()
    return render(
        request,
        "energy/dashboard.html",
        {
            "building_count": building_count,
            "last_building": last_building,
        },
    )


def building_create(request):
    result = None

    if request.method == "POST":
        form = BuildingForm(request.POST)
        if form.is_valid():
            building = form.save()

            # Berechnung durchführen
            result = calc_heating_demand(building)

            # Ergebnisse ins Modell schreiben
            building.result_Q_T = result["Q_T"]
            building.result_Q_V = result["Q_V"]
            building.result_Q_I = result["Q_I"]
            building.result_Q_S = result["Q_S"]
            building.result_Q_h = result["Q_h"]

            building.result_H_T = result["H_T"]
            building.result_H_V = result["H_V"]

            building.result_floor_area = result["floor_area"]
            building.result_roof_area = result["roof_area"]
            building.result_opaque_wall_area = result["opaque_wall_area"]
            building.result_window_area = result["window_area"]

            building.result_Q_S_n = result["Q_S_n"]
            building.result_Q_S_e = result["Q_S_e"]
            building.result_Q_S_s = result["Q_S_s"]
            building.result_Q_S_w = result["Q_S_w"]

            building.result_Q_PV_total = result["Q_PV_total"]
            building.result_Q_PV_on = result["Q_PV_on"]
            building.result_Q_PV_off = result["Q_PV_off"]

            building.save()

            # >>> HIER wird zur Gebäudeliste umgeleitet <<<
            return redirect("building_list")

    else:
        form = BuildingForm()

    context = {
        "form": form,
        "result": result,
    }
    return render(request, "energy/building_form.html", {"form": form, "result": result})


def building_list(request):
    # erlaubte Sortierfelder
    order = request.GET.get("order", "-id")
    allowed_orders = {
        "name": "name",
        "name_desc": "-name",
        "area": "result_floor_area",
        "area_desc": "-result_floor_area",
        "q_h": "result_Q_h",
        "q_h_desc": "-result_Q_h",
        "id": "-id",
    }
    order_by = allowed_orders.get(order, "-id")

    buildings = Building.objects.all().order_by(order_by)
    context = {
        "buildings": buildings,
        "current_order": order,
    }
    return render(request, "energy/building_list.html", context)


def building_delete(request, pk):
    building = get_object_or_404(Building, pk=pk)
    if request.method == "POST":
        building.delete()
        return redirect("building_list")
    # falls jemand per GET draufgeht: einfach zurück zur Liste
    return redirect("building_list")


def building_delete_all(request):
    if request.method == "POST":
        Building.objects.all().delete()
        return redirect("building_list")
    return redirect("building_list")


def building_edit(request, pk):
    building = get_object_or_404(Building, pk=pk)

    if request.method == "POST":
        form = BuildingForm(request.POST, instance=building)
        if form.is_valid():
            building = form.save()

            # Berechnung erneut durchführen
            result = calc_heating_demand(building)

            building.result_Q_T = result["Q_T"]
            building.result_Q_V = result["Q_V"]
            building.result_Q_I = result["Q_I"]
            building.result_Q_S = result["Q_S"]
            building.result_Q_h = result["Q_h"]

            building.result_H_T = result["H_T"]
            building.result_H_V = result["H_V"]

            building.result_floor_area = result["floor_area"]
            building.result_roof_area = result["roof_area"]
            building.result_opaque_wall_area = result["opaque_wall_area"]
            building.result_window_area = result["window_area"]

            building.result_Q_S_n = result["Q_S_n"]
            building.result_Q_S_e = result["Q_S_e"]
            building.result_Q_S_s = result["Q_S_s"]
            building.result_Q_S_w = result["Q_S_w"]

            building.result_Q_PV_total = result["Q_PV_total"]
            building.result_Q_PV_on = result["Q_PV_on"]
            building.result_Q_PV_off = result["Q_PV_off"]

            building.save()

            return redirect("building_detail", pk=building.pk)
    else:
        form = BuildingForm(instance=building)

    return render(
        request,
        "energy/building_form.html",
        {"form": form, "edit_mode": True, "building": building},
    )


def building_export_csv(request):
    """
    Export aller Gebäude als CSV.
    """
    buildings = Building.objects.all().order_by("id")

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="buildings_export.csv"'

    writer = csv.writer(response, delimiter=";")

    # Kopfzeile
    writer.writerow([
        "ID",
        "Name",
        "Grundfläche [m²]",
        "Q_h [kWh/a]",
        "Q_T [kWh/a]",
        "Q_V [kWh/a]",
        "Q_I [kWh/a]",
        "Q_S [kWh/a]",
        "PV gesamt [kWh/a]",
        "PV Eigenverbrauch [kWh/a]",
        "PV Überschuss [kWh/a]",
    ])

    # Datenzeilen
    for b in buildings:
        writer.writerow([
            b.id,
            b.name,
            b.result_floor_area or "",
            b.result_Q_h or "",
            b.result_Q_T or "",
            b.result_Q_V or "",
            b.result_Q_I or "",
            b.result_Q_S or "",
            b.result_Q_PV_total or "",
            b.result_Q_PV_on or "",
            b.result_Q_PV_off or "",
        ])

    return response


def building_export_xlsx(request):
    """
    Export aller Gebäude als Excel-Datei (XLSX).
    """
    buildings = Building.objects.all().order_by("id")

    wb = Workbook()
    ws = wb.active
    ws.title = "Gebäude"

    # Kopfzeile
    headers = [
        "ID",
        "Name",
        "Grundfläche [m²]",
        "Q_h [kWh/a]",
        "Q_T [kWh/a]",
        "Q_V [kWh/a]",
        "Q_I [kWh/a]",
        "Q_S [kWh/a]",
        "PV gesamt [kWh/a]",
        "PV Eigenverb [kWh/a]",
        "PV Überschuss [kWh/a]",
    ]
    ws.append(headers)

    # Kopfzeile fett machen
    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font

    # Datenzeilen
    for b in buildings:
        ws.append([
            b.id,
            b.name,
            b.result_floor_area or "",
            b.result_Q_h or "",
            b.result_Q_T or "",
            b.result_Q_V or "",
            b.result_Q_I or "",
            b.result_Q_S or "",
            b.result_Q_PV_total or "",
            b.result_Q_PV_on or "",
            b.result_Q_PV_off or "",
        ])

    # In Memory-Stream schreiben
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    # HTTP-Response
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="buildings_export.xlsx"'
    return response


def building_export_pdf(request):
    """
    Export aller Gebäude als übersichtliche PDF-Tabelle.
    """
    buildings = Building.objects.all().order_by("id")

    # PDF in Memory erzeugen
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    elements = []

    # Titel
    elements.append(Paragraph("Gebäude-Export – Energiebilanz", styles["Title"]))
    elements.append(Spacer(1, 8))

    # kleine Zusammenfassung
    info_text = f"Anzahl Gebäude: {buildings.count()}"
    elements.append(Paragraph(info_text, styles["Normal"]))
    elements.append(Spacer(1, 12))

    # Hilfsfunktion für Zahlenformat
    def fmt(val, decimals=1):
        if val is None:
            return "-"
        return f"{val:.{decimals}f}"

    # Tabellendaten
    data = [
        ["ID", "Name", "Grundfl. [m²]", "Q_h [kWh/a]", "PV on [kWh/a]", "PV off [kWh/a]"]
    ]

    for b in buildings:
        data.append([
            str(b.id),
            b.name,
            fmt(b.result_floor_area, 1),
            fmt(b.result_Q_h, 1),
            fmt(b.result_Q_PV_on, 1),
            fmt(b.result_Q_PV_off, 1),
        ])

    # Tabelle mit Styling
    table = Table(data, repeatRows=1, hAlign="LEFT")

    table.setStyle(TableStyle([
        # Header
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9ecef")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
        ("TOPPADDING", (0, 0), (-1, 0), 6),

        # Körper
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
        ("TOPPADDING", (0, 1), (-1, -1), 4),

        # Gitterlinien
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),

        # Ausrichtung: ID + Zahlen rechts, Name links
        ("ALIGN", (0, 1), (0, -1), "RIGHT"),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),

        # wechselnde Zeilenfarben im Body
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#f8f9fa")]),
    ]))

    elements.append(table)

    # PDF aufbauen
    doc.build(elements)

    # Response zurückgeben
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf"
    )
    response["Content-Disposition"] = 'attachment; filename=\"buildings_export.pdf\"'
    return response


def building_result_pdf(request, pk):
    """
    PDF-Bericht für EIN Gebäude inkl. Tabelle und einfachem Balkendiagramm.
    """
    building = get_object_or_404(Building, pk=pk)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    elements = []

    # Titel
    title = f"Energiebericht – {building.name}"
    elements.append(Paragraph(title, styles["Title"]))
    elements.append(Spacer(1, 10))

    # Kurzinfo
    info_lines = [
        f"Gebäude-ID: {building.id}",
        f"Grundfläche: {building.result_floor_area or '-'} m²",
        f"Raum-Solltemperatur: {building.setpoint_temp or '-'} °C",
    ]
    for line in info_lines:
        elements.append(Paragraph(line, styles["Normal"]))
    elements.append(Spacer(1, 12))

    # Hilfsfunktion Zahlenformat
    def fmt(val, decimals=1):
        if val is None:
            return "-"
        return f"{val:.{decimals}f}"

    # Ergebnis-Tabelle
    data = [
        ["Größe", "Wert"],
        ["Heizwärmebedarf Q_h [kWh/a]", fmt(building.result_Q_h, 1)],
        ["Transmissionsverluste Q_T [kWh/a]", fmt(building.result_Q_T, 1)],
        ["Lüftungsverluste Q_V [kWh/a]", fmt(building.result_Q_V, 1)],
        ["Interne Gewinne Q_I [kWh/a]", fmt(building.result_Q_I, 1)],
        ["Solare Gewinne Q_S [kWh/a]", fmt(building.result_Q_S, 1)],
        ["PV-Gesamt [kWh/a]", fmt(building.result_Q_PV_total, 1)],
        ["PV-Eigenverbrauch [kWh/a]", fmt(building.result_Q_PV_on, 1)],
        ["PV-Überschuss [kWh/a]", fmt(building.result_Q_PV_off, 1)],
    ]

    table = Table(data, hAlign="LEFT", colWidths=[70 * mm, 60 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9ecef")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(Paragraph("Ergebnisse – Übersicht", styles["Heading2"]))
    elements.append(table)
    elements.append(Spacer(1, 16))

    # Balkendiagramm Energieflüsse
    elements.append(Paragraph("Energieflüsse (jährlich)", styles["Heading2"]))
    elements.append(Spacer(1, 4))

    flows = [
        ("Q_T", building.result_Q_T or 0),
        ("Q_V", building.result_Q_V or 0),
        ("Q_I", building.result_Q_I or 0),
        ("Q_S", building.result_Q_S or 0),
        ("Q_h", building.result_Q_h or 0),
    ]

    data_values = [[v for (_, v) in flows]]
    labels = [lbl for (lbl, _) in flows]

    drawing = Drawing(400, 200)
    bc = VerticalBarChart()
    bc.x = 40
    bc.y = 40
    bc.height = 120
    bc.width = 320
    bc.data = data_values
    bc.categoryAxis.categoryNames = labels
    bc.categoryAxis.labels.boxAnchor = "n"
    bc.valueAxis.valueMin = 0
    bc.barWidth = 20
    bc.groupSpacing = 10
    bc.barSpacing = 5
    bc.bars[0].fillColor = colors.HexColor("#0d6efd")

    drawing.add(bc)
    elements.append(drawing)

    # PDF bauen
    doc.build(elements)

    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="building_{building.id}_bericht.pdf"'
    )
    return response


def building_detail(request, pk):
    """
    Detailansicht eines Gebäudes mit gespeichertem Ergebnis.
    (Template verwenden wir später, wenn du soweit bist.)
    """
    building = get_object_or_404(Building, pk=pk)
    return render(request, "energy/building_detail.html", {"building": building})
